"""The Word files (decision 2a).

Review copy  script/review_v<N>.docx: every spoken line is one paragraph; its material type, its sources (outlet + date,
             never a link), the newspaper headline of news lines, pronunciation, and every warning are a margin comment.
             A hidden bookmark carries each line's id, so your edits come back to the right line. You edit the text in
             Word as usual; to change a line's material type, start it with the type in brackets, e.g. «[خريطة] …».
             A new paragraph you add is a new line (no source: commentary, so it may hold no number). Deleting a
             paragraph deletes the line. Then run  bin/monteer script "<project>"  again: the edits are imported, checked,
             and a new review copy v<N+1> is written.
Sources file script/sources_v<N>.md (INTERNAL): every card with outlet, date, link and the checked passage.
Recording copy script/script.docx: made only by `script approve`: title, part headers and stage directions in
             parentheses, the spoken lines, foreign names glossed as "(Latin) (النطق: …)" the first time. No links, no
             markers, no comments. It parses (cut/script_text.py) into exactly the approved lines.
"""
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..core.io import atomic, load_json, save_json, write_text
from .common import S, TYPE_FROM_AR, TYPES, cards, latest_version, now, sha

AUTHOR = "مونتير"
TYPE_TAG = re.compile(r"^\s*\[([^\]]{2,20})\]\s*")


def _rtl(par, size=14, bold=False, color=None):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:bidi")
    pPr.append(b)
    for r in par.runs:
        rPr = r._r.get_or_add_rPr()
        rPr.append(OxmlElement("w:rtl"))
        fonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), "Arial")
        rPr.insert(0, fonts)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(size * 2))
        rPr.append(szcs)
        r.font.size = Pt(size)
        r.bold = bold
        if bold:
            rPr.append(OxmlElement("w:bCs"))
        if color:
            r.font.color.rgb = RGBColor(*color)


_bm = [0]


def _bookmark(par, name):
    _bm[0] += 1
    s, e = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    s.set(qn("w:id"), str(_bm[0]))
    s.set(qn("w:name"), name)
    e.set(qn("w:id"), str(_bm[0]))
    p = par._p
    first = p.find(qn("w:r"))
    if first is not None:
        first.addprevious(s)
    else:
        p.append(s)
    p.append(e)


def _para(doc, text, mark, size=14, bold=False, color=None):
    par = doc.add_paragraph()
    par.add_run(text)
    _rtl(par, size, bold, color)
    _bookmark(par, mark)
    return par


def _comment(doc, par, lines):
    c = doc.add_comment(par.runs, text=lines[0], author=AUTHOR, initials="M")
    for ln in lines[1:]:
        c.add_paragraph(ln)


def ar_date(d):
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", d or "")
    months = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
    return f"{int(m.group(3))} {months[int(m.group(2)) - 1]} {m.group(1)}" if m else (d or "بلا تاريخ")


def ar_period(d):
    """2024 -> 2024 · 2026-09 -> سبتمبر 2026 · 2026-09-04 -> 4 سبتمبر 2026 · 1970–2024 -> 1970–2024"""
    months = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
    def one(x):
        m = re.fullmatch(r"(\d{4})M(\d{2})", x)          # the World Bank's monthly periods: 2026M07
        if m:
            x = f"{m.group(1)}-{m.group(2)}"
        m = re.fullmatch(r"(\d{4})Q([1-4])", x)           # and quarterly ones: 2026Q1
        if m:
            return f"الربع {['الأول', 'الثاني', 'الثالث', 'الرابع'][int(m.group(2)) - 1]} {m.group(1)}"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", x):
            return ar_date(x)
        if re.fullmatch(r"\d{4}-\d{2}", x):
            return f"{months[int(x[5:7]) - 1]} {x[:4]}"
        return x
    return "–".join(one(x) for x in (d or "").split("–")) if d else ""


def card_when(c):
    """«البنك الدولي – بيانات 1974» for data (the data's own period, report decision 5a); the publication date otherwise"""
    if c.get("source_kind") == "data" or c.get("period"):   # a number dated by its data's period (must-do g)
        per = c.get("period") or c.get("series_period")
        return f"بيانات {ar_period(per)}" if per else "بيانات بلا فترة"
    return ar_date(c.get("pub_date"))


def comment_lines(x, cs, probs):
    out = [f"النوع: {TYPES.get(x['type'], x['type'])}" + (" — تعليق/تحليل (بلا مصدر)" if x.get("commentary") else "")]
    for cid in x.get("cards", []):
        c = cs.get(cid)
        if not c:
            out.append(f"[{cid}] ؟ (بطاقة غير موجودة)")
            continue
        s = f"[{cid}] {c['outlet']} – {card_when(c)}"
        if c["kind"] == "quote":
            s += f" — اقتباس: {c.get('speaker', '')}"
            if c.get("quote_original") and c["quote_original"].strip() != c.get("quote_ar", "").strip():
                s += f" — النص الأصلي: «{c['quote_original']}» (ترجمتنا)"
        if c.get("event_date") and c.get("source_kind") != "data":
            s += f" — تاريخ الحدث: {c['event_date']}"
        out.append(s)
    if x.get("news"):
        n = x["news"]
        out.append(f"الجريدة: «{n.get('headline', '')}» — {n.get('context', '')} — تمييز: {'، '.join(n.get('highlight', []))}")
    for nm in x.get("names", []):
        out.append(f"النطق: {nm.get('ar', '')} = {nm.get('latin', '')} — {nm.get('say', '')}")
    for pr in probs:
        mark = {"error": "❌ خطأ: ", "warning": "⚠ تنبيه: ", "label": "🤖 ", "info": "ℹ "}[pr["level"]]
        out.append(mark + pr["text"])
    if x.get("edited"):
        out.append("✎ عُدّل يدويًا")
    return out


def write_review(p, v=None):
    v = v or latest_version(p)
    dr = load_json(S(p, "drafts", f"v{v}.json"))
    res = load_json(S(p, "drafts", f"v{v}.checks.json"))
    cs = {c["id"]: c for c in cards(p, only_ok=False)}
    _bm[0] = 0
    doc = Document()
    sec = doc.sections[0]
    sec._sectPr.append(OxmlElement("w:bidi"))
    local = res["script"]["local_lines"]
    makers = sorted({x.get("made_by", "") for x in dr["items"] if x["kind"] == "line" and x.get("made_by")})
    _para(doc, f"نسخة المراجعة — الإصدار {v} — {now()[:16]}", "META_1", 10, color=(90, 90, 90))
    if local:
        _para(doc, f"⚠ مسودة من نموذج محلي مجاني – جودة أقل – راجِع كل سطر ({local} سطر). الموافقة تحتاج  --accept-local", "META_2", 13, True, (192, 0, 0))
    _para(doc, "كتبه: " + "، ".join(makers or ["?"]) + f" · الأخطاء: {res['totals']['errors']} · التنبيهات: {res['totals']['warnings']} · "
               f"الكلمات: {res['script']['words']} من {res['script']['target_words']}", "META_3", 10, color=(90, 90, 90))
    _para(doc, "للتعديل: عدّل النص مباشرة. لتغيير نوع المادة ابدأ السطر بالنوع بين قوسين، مثل [خريطة]. السطر الجديد بلا مصدر = تعليق "
               "(بلا أرقام). المصادر وأنواع المواد في الهامش؛ الروابط في الملف الداخلي فقط.", "META_4", 10, color=(90, 90, 90))
    for w in res["script"]["warnings"]:
        _para(doc, "⚠ " + w, f"META_W{_bm[0]}", 10, color=(160, 90, 0))
    _para(doc, dr["title"], "TITLE", 18, True)
    for k, x in enumerate(dr["items"]):
        if x["kind"] == "line":
            par = _para(doc, x["text"], x["id"])
            _comment(doc, par, comment_lines(x, cs, res["lines"].get(x["id"], [])))
        else:
            _para(doc, x["text"], f"{x['kind'].upper()}_{k}", 12, x["kind"] == "part", (60, 60, 140))
    out = S(p, f"review_v{v}.docx")
    with atomic(out) as tmp:
        doc.save(str(tmp))
    save_json(S(p, "drafts", f"v{v}.review.json"), {"file": out.name, "sha": sha(out.read_bytes()), "written_at": now()})
    write_sources(p, v, dr, cs)
    print(f"review copy: {out}")
    return out


def write_sources(p, v, dr, cs):
    used = {}
    for x in dr["items"]:
        if x["kind"] == "line":
            for c in x.get("cards", []):
                used.setdefault(c, []).append(x["id"])
    L = [f"# Sources — draft v{v} (INTERNAL: links stay here, never on screen or in the description)", "",
         f"{len(used)} cards used by the draft; {sum(1 for c in cs.values() if c['status'] == 'ok')} cards passed their checks; "
         f"{sum(1 for c in cs.values() if c['status'] != 'ok')} were dropped.", ""]
    for cid, c in sorted(cs.items()):
        if cid not in used:
            continue
        when = (f"data {c.get('period') or c.get('series_period') or '?'}" if c["source_kind"] == "data" or c.get("period") else (c.get("pub_date") or "no date"))
        L += [f"## {cid} — {c['outlet']} – {when} ({c['source_kind']}, {c['kind']})", "",
              f"- Used by: {', '.join(used[cid])}", f"- Title: {c['title']}", f"- Link: {c['url']}",
              f"- Card: {c['statement']}", f"- Checked passage: «{c['passage']}»"]
        if c["kind"] == "quote":
            L.append(f"- Quote: {c.get('speaker')}: «{c.get('quote_original')}» → «{c.get('quote_ar')}»" + (" — DISPUTED" if c.get("disputed") else ""))
        for f in c.get("flags", []):
            L.append(f"- ⚠ {f}")
        if c["source_kind"] == "data":
            L.append(f"- Series covers: {c.get('series_period') or '?'} · fetched: {c.get('fetched_at') or '?'}")
        L.append(f"- Card made by: {c.get('made_by', '?')}")
        L.append("")
    L += ["## Cards not used", ""] + [f"- {cid}: {c['outlet']} – {card_when(c)}: {c['statement']}" for cid, c in sorted(cs.items())
                                      if cid not in used and c["status"] == "ok"]
    L += ["", "## Cards dropped by the checks", ""] + [f"- {cid} ({c['outlet']}): {'; '.join(c['problems'])}" for cid, c in sorted(cs.items())
                                                       if c["status"] != "ok"]
    write_text(S(p, f"sources_v{v}.md"), "\n".join(L) + "\n")


# ------------------------------------------------------------------------------------------------ import your edits
def _marks(par):
    return [m for m in par._p.xpath(".//w:bookmarkStart/@w:name") if not m.startswith("_")]


def guess_type(text):
    from .text import numbers
    ns = numbers(text)
    if any(1800 <= n["value"] <= 2015 and float(n["value"]).is_integer() for n in ns):
        return "archive"
    return "numbers" if ns else "stock"


def read_edits(p):
    """the latest review copy, if you changed it since it was written: a new raw draft v<N+1>, else None"""
    v = latest_version(p)
    if not v:
        return None
    meta_f = S(p, "drafts", f"v{v}.review.json")
    f = S(p, f"review_v{v}.docx")
    if not (meta_f.exists() and f.exists()) or load_json(meta_f)["sha"] == sha(f.read_bytes()):
        return None
    dr = load_json(S(p, "drafts", f"v{v}.json"))
    old = {x["id"]: x for x in dr["items"] if x["kind"] == "line"}
    doc = Document(str(f))
    items, title = [], dr["title"]
    nxt = max([int(i[1:]) for i in old] + [0])
    changed = []
    for par in doc.paragraphs:
        text = par.text.strip()
        marks = _marks(par)
        if any(m.startswith("META") for m in marks):
            continue
        if "TITLE" in marks:
            title = text or title
            continue
        lid = next((m for m in marks if re.fullmatch(r"L\d{3,}", m)), None)
        if not text:
            continue
        kind = next((mm.group(1).lower() for m in marks if (mm := re.fullmatch(r"(PART|STAGE|PROMO_START|PROMO_END)_\d+", m))), None)
        if lid and lid in old:
            x = dict(old[lid])
            m = TYPE_TAG.match(text)
            if m and m.group(1).strip() in TYPE_FROM_AR:
                x["type"] = TYPE_FROM_AR[m.group(1).strip()]
                text = text[m.end():].strip()
                changed.append(f"{lid}: type -> {x['type']}")
            if text != x["text"]:
                x.update(text=text, edited=True)
                changed.append(f"{lid}: text edited")
            items.append(x)
            old.pop(lid)
            continue
        if kind or (text.startswith("(") and text.endswith(")")):
            items.append({"kind": kind or "stage", "text": text})
            continue
        nxt += 1
        m = TYPE_TAG.match(text)
        t = TYPE_FROM_AR.get(m.group(1).strip()) if m else None
        if m:
            text = text[m.end():].strip()
        x = {"kind": "line", "id": f"L{nxt:03d}", "text": text, "type": t or guess_type(text), "cards": [], "commentary": True,
             "names": [], "news": None, "made_by": "edited by hand", "edited": True, "new": True}
        items.append(x)
        changed.append(f"{x['id']}: new line")
    for lid in old:
        changed.append(f"{lid}: deleted")
    nd = dict(dr, version=v + 1, title=title, items=items, made_at=now(), from_review=f"review_v{v}.docx", edits=changed)
    for k in ("verified_at", "revise_rounds"):
        nd.pop(k, None)
    return nd, changed
