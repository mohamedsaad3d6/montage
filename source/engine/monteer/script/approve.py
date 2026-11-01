"""Step import (your Word edits come back) and the command  bin/monteer script approve "<project>" --name "<who>".

Approval (decision 11a) is refused while: the latest review copy has edits not yet imported and checked; any error is
left; any line was written by the local model and --accept-local is not given. When the research is more
than 48 hours old, every news and official card's page is fetched again and its passage must still be there.
Approval writes:
  approved.json                    who, when, version, the text's fingerprint, the brains used
  script.docx                      the recording copy, checked to parse (cut/script_text.py) into exactly the lines
  lines.json                       every spoken line: id, part, text, sentences (as the cut reads them), type, cards
  news_items.json                  documented news for the montage's newspaper graphic (decision 7 shape)
  data_series.json                 the data series behind the cited data cards
  project.json inputs.script       -> script/script.docx
Any later edit (a new review copy imported) cancels the approval until checked and approved again.
"""
import json
import os
import re
import time

from docx import Document

from ..core import brain
from ..core.io import atomic, load_json, save_json
from ..cut import script_text
from . import extract, fetch
from .common import S, TYPES, cards, channel_dir, latest_version, now, sha, style_dir
from .research import request
from .review import _rtl, ar_date, read_edits
from .text import contains, norm, numbers, number_found

SOURCING_SCHEMA = {"type": "object", "required": ["lines"], "properties": {"lines": {"type": "array", "items": {
    "type": "object", "required": ["id", "cards"], "properties": {"id": {"type": "string"},
                                                                 "cards": {"type": "array", "items": {"type": "string"}}}}}}}


def cancel_approval(p, why):
    a = S(p, "approved.json")
    if a.exists():
        rec = load_json(a)
        dst = S(p, f"approved_cancelled_{time.strftime('%Y%m%d-%H%M%S')}.json")
        os.replace(a, dst)
        save_json(dst, dict(rec, cancelled_at=now(), cancelled_why=why))     # the panel shows when and why
        print(f"APPROVAL CANCELLED: {why}. Check and approve again.")


def step_import(p):
    from .panel_edits import read as panel_read       # the owner's line edits made in the panel (monteer/script/panel_edits.py)
    got = read_edits(p) or panel_read(p)
    if not got:
        print("no edits in the latest review copy")
        return
    nd, changed = got
    if not changed:                     # saved again without a change to the text: nothing to check, nothing cancelled
        v = nd["version"] - 1
        f = S(p, f"review_v{v}.docx")
        meta = load_json(S(p, "drafts", f"v{v}.review.json"))
        save_json(S(p, "drafts", f"v{v}.review.json"), dict(meta, sha=sha(f.read_bytes()), resaved_at=now()))
        print(f"{f.name} was saved again without changes to the text: nothing to import")
        return
    print(f"imported your edits from {nd['from_review']}: " + "; ".join(changed))
    cancel_approval(p, "the script was edited in the panel after approval" if nd.get("from_panel") else "the review copy was edited after approval")
    cs = {c["id"]: c for c in cards(p)}
    need = [x for x in nd["items"] if x["kind"] == "line" and x.get("edited")
            and any(not number_found(n, [m for cid in x["cards"] if cid in cs for f in ("statement", "passage", "quote_ar")
                                         for m in numbers(cs[cid].get(f, ""))]) for n in numbers(x["text"]))]
    if need:
        instructions = ("The owner edited these script lines by hand. Some now hold numbers or facts that the cards they cite "
                        "do not cover. For each line, list the ids of cards (only from inputs.cards) that support EVERYTHING "
                        "factual in it; an empty list if no card supports it. Never invent a card. Answer with JSON only.")
        inputs = {"lines": [{"id": x["id"], "text": x["text"], "cards_now": x["cards"]} for x in need],
                  "cards": [{"id": c["id"], "statement": c["statement"], "outlet": c["outlet"], "date": c.get("pub_date", "")} for c in cs.values()]}
        d = brain.decide(p, "script", f"sourcing_v{nd['version']}", request("find cards for edited lines", instructions, inputs, SOURCING_SCHEMA, 1500,
                                                                                     local_ok=True))
        found = {x["id"]: [c for c in x["cards"] if c in cs] for x in d["lines"]}
        for x in need:
            if found.get(x["id"]):
                x["cards"] = sorted(set(x["cards"]) | set(found[x["id"]]))
                x["commentary"] = False
    save_json(S(p, "drafts", f"v{nd['version']}.raw.json"), nd)


# ------------------------------------------------------------------------------------------------ approve
def gloss(text, names, done):
    """the first time a foreign name appears: «كينز (Keynes) (النطق: كَيْنْز)»"""
    for nm in names:
        key = norm(nm.get("ar", ""))
        if not key or key in done or nm.get("ar", "") not in text:
            continue
        add = f" ({nm['latin']})" if nm.get("latin") and re.match(r"[A-Za-z]", nm["latin"]) else ""
        if nm.get("say"):
            add += f" (النطق: {nm['say']})"
        text = text.replace(nm["ar"], nm["ar"] + add, 1)
        done.add(key)
    return text


def recording_copy(p, dr):
    doc = Document()
    paras = [dr["title"]]
    done = set()
    for x in dr["items"]:
        paras.append(gloss(x["text"], x.get("names", []), done) if x["kind"] == "line" else x["text"])
    for t in paras:
        par = doc.add_paragraph()
        par.add_run(t)
        _rtl(par, 18 if t is paras[0] else 14, t is paras[0])
    out = S(p, "script.docx")
    with atomic(out) as tmp:
        doc.save(str(tmp))
    return out


def parse_check(path, dr):
    """the recording copy, read the way the cut reads scripts, gives exactly the approved lines' words"""
    paras = script_text.docx_paragraphs(path)
    got = []
    for para in paras[1:]:
        if script_text.is_stage_direction(para):
            continue
        for s in script_text.split_sentences(script_text.strip_glosses(para)):
            got.append(s)
    want, sent_map = [], {}
    for x in dr["items"]:
        if x["kind"] != "line":
            continue
        ss = script_text.split_sentences(x["text"])
        sent_map[x["id"]] = list(range(len(want), len(want) + len(ss)))
        want += ss
    gw = [w for s in got for w in norm(s).split()]
    ww = [w for s in want for w in norm(s).split()]
    return gw == ww, len(got), len(want), sent_map, want


def word_diff(path, dr):
    """what the recording copy has that the approved lines don't, and the other way -> (extra words, missing words), as
    written (the comparison is parse_check's: the normalized words)"""
    import difflib
    paras = script_text.docx_paragraphs(path)
    got = [s for para in paras[1:] if not script_text.is_stage_direction(para)
           for s in script_text.split_sentences(script_text.strip_glosses(para))]
    want = [s for x in dr["items"] if x["kind"] == "line" for s in script_text.split_sentences(x["text"])]

    def toks(ss):
        return [(norm(w), w) for s in ss for w in s.split() if norm(w).strip()]
    g, w = toks(got), toks(want)
    extra, missing = [], []
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(a=[x[0] for x in w], b=[x[0] for x in g], autojunk=False).get_opcodes():
        if op in ("insert", "replace"):
            extra.append(" ".join(x[1] for x in g[b1:b2]))
        if op in ("delete", "replace"):
            missing.append(" ".join(x[1] for x in w[a1:a2]))
    return extra, missing


def not_parsed_message(extra, missing):
    """the refusal in Arabic, naming the words that differ (the panel shows it as it is)"""
    parts = []
    if extra:
        parts.append("فيها كلمات مش موجودة في السطور المعتمدة: " + "، ".join(f"«{x}»" for x in extra[:6]))
    if missing:
        parts.append("ناقصها كلمات من السطور المعتمدة: " + "، ".join(f"«{x}»" for x in missing[:6]))
    return ("ماتعتمدش: نسخة التسجيل (script.docx) مش مطابقة للسطور المعتمدة كلمة بكلمة؛ "
            + ("؛ ".join(parts) if parts else "الاختلاف في ترتيب الكلمات") + ".")


def refetch(p, cs, hours=48):
    """re-read the pages of news/official cards when the research is older than `hours`"""
    items = load_json(S(p, "research", "items.json"))
    t0 = time.mktime(time.strptime(items["gathered_at"], "%Y-%m-%d %H:%M:%S"))
    age = (time.time() - t0) / 3600
    if age < hours:
        return {"age_hours": round(age, 1), "refetched": 0, "problems": []}
    probs, pages = [], {}
    for c in cs:
        if c["source_kind"] in ("news", "official") and c.get("url"):
            pages.setdefault(c["url"], []).append(c)
    for url, group in pages.items():             # each page once, however many cards come from it
        r = fetch.get(url, max_age=0)
        if not r["ok"]:
            probs += [f"{c['id']} ({c['outlet']}): الصفحة مابقتش بتفتح ({r.get('code')} {r.get('why') or r.get('error') or ''})" for c in group]
            continue
        text = extract.article(r["body"], url)["text"]
        probs += [f"{c['id']} ({c['outlet']}): الصفحة اتغيرت والفقرة مابقتش موجودة فيها (اتصححت؟)"
                  for c in group if not contains(text, c["passage"])]
    return {"age_hours": round(age, 1), "refetched_pages": len(pages), "cards": sum(len(g) for g in pages.values()), "problems": probs}


def approve(p, name, accept_local=False):
    v = latest_version(p)
    if not v:
        raise SystemExit("لسه مفيش مسودة اتراجعت: شغّل كاتب السكريبت الأول (bin/monteer script \"<project>\")")
    rv = S(p, "drafts", f"v{v}.review.json")
    if not rv.exists() or load_json(rv)["sha"] != sha(S(p, f"review_v{v}.docx").read_bytes()):
        raise SystemExit(f"نسخة المراجعة review_v{v}.docx فيها تعديلات لسه مادخلتش ومااتراجعتش: شغّل كاتب السكريبت الأول "
                         f"(bin/monteer script \"{p.folder}\")")
    dr = load_json(S(p, "drafts", f"v{v}.json"))
    res = load_json(S(p, "drafts", f"v{v}.checks.json"))
    if res["totals"]["errors"]:
        raise SystemExit(f"ماتعتمدش: لسه فيه {res['totals']['errors']} أخطاء (شوفها في drafts/v{v}.checks.md أو في هامش نسخة المراجعة)")
    if res["script"]["local_lines"] and not accept_local:
        raise SystemExit(f"ماتعتمدش: {res['script']['local_lines']} سطر كتبهم النموذج المحلي المجاني؛ اقراهم كويس، وبعدين "
                         "اعتمد مع «قريت سطور النموذج المحلي كويس» (أو --accept-local)")
    used = sorted({c for x in dr["items"] if x["kind"] == "line" for c in x.get("cards", [])})
    allc = {c["id"]: c for c in cards(p, only_ok=False)}
    fresh = refetch(p, [allc[c] for c in used if c in allc])
    if fresh["problems"]:
        raise SystemExit("ماتعتمدش: معلومات من الأخبار اتغيرت من ساعة ما اتجمعت:\n  " + "\n  ".join(fresh["problems"]))
    out = recording_copy(p, dr)
    ok, ng, nw, sent_map, want = parse_check(out, dr)
    if not ok:
        raise SystemExit(not_parsed_message(*word_diff(out, dr)))
    lines = []
    news = []
    for x in dr["items"]:
        if x["kind"] != "line":
            continue
        ln = {"id": x["id"], "part": x.get("part"), "text": x["text"], "sentences": sent_map[x["id"]],
              "type": "promo" if x.get("promo") else x["type"], "type_ar": TYPES.get(x["type"], "إعلان"), "cards": x.get("cards", []),
              "commentary": x.get("commentary", False), "made_by": x.get("made_by", "")}
        if x["type"] == "news" and x.get("news"):
            c = next((allc[c] for c in x["cards"] if c in allc and allc[c]["source_kind"] in ("news", "official")), None)
            ni = {"id": f"NI{len(news) + 1:02d}", "line_ids": [x["id"]], "headline": x["news"].get("headline", ""),
                  "context": x["news"].get("context", ""), "highlight": x["news"].get("highlight", []),
                  "source_line": f"المصدر: {c['outlet']} – {ar_date(c.get('pub_date'))}" if c else "",
                  "outlet": c["outlet"] if c else "", "date": c.get("pub_date", "") if c else "", "card_ids": x["cards"],
                  "url_internal": c["url"] if c else "", "verified": True}
            news.append(ni)
            ln["news_item"] = ni["id"]
        lines.append(ln)
    items = {i["id"]: i for i in load_json(S(p, "research", "items.json"))["items"]}
    data = []
    for cid in used:
        c = allc.get(cid)
        if c and c["source_kind"] == "data":
            it = items.get(c["item_id"], {})
            txt = S(p, "research", "texts", f"{c['item_id']}.txt")
            data.append({"card": cid, "item": c["item_id"], "source": c["outlet"], "period": c.get("period") or c.get("series_period", ""),
                         "series_period": c.get("series_period", ""), "request": it.get("data_request"), "url_internal": c["url"],
                         "series_text": txt.read_text(encoding="utf-8") if txt.exists() else ""})
    save_json(S(p, "lines.json"), {"version": v, "lines": lines, "sentences": want})
    save_json(S(p, "news_items.json"), {"version": v, "items": news,
                                        "note": "decision 7: our own newspaper graphic; the headline is a real published fact in our words; "
                                                "the outlet appears in plain text only; url_internal never goes on screen"})
    save_json(S(p, "data_series.json"), {"version": v, "series": data})
    brains = sorted({x.get("made_by", "") for x in dr["items"] if x["kind"] == "line"})
    save_json(S(p, "approved.json"), {"by": name, "at": now(), "version": v, "sha": sha([x["text"] for x in dr["items"]]),
                                      "brains": brains, "accepted_local": bool(res["script"]["local_lines"]), "refetch": fresh,
                                      "recording_copy": out.name})
    cfg = json.loads((p.folder / "project.json").read_text(encoding="utf-8"))
    cfg.setdefault("inputs", {})["script"] = "script/script.docx"
    save_json(p.folder / "project.json", cfg)
    learn_from_edits(p, dr)
    print(f"approved by {name}: v{v}, {len(lines)} lines, {len(news)} news items, {len(data)} data series -> {out}")
    return out


def learn_from_edits(p, final):
    """the pairs (first draft -> approved) of lines you edited: the next drafts read them (style/edits.json)"""
    first = load_json(S(p, "drafts", "v1.json"))
    before = {x["id"]: x["text"] for x in first["items"] if x["kind"] == "line"}
    pairs = [{"draft": before[x["id"]], "approved": x["text"], "project": p.name} for x in final["items"]
             if x["kind"] == "line" and x["id"] in before and norm(before[x["id"]]) != norm(x["text"])]
    if not pairs:
        return
    f = style_dir(channel_dir(p)) / "edits.json"
    d = load_json(f) if f.exists() else {"pairs": []}
    d["pairs"] = (d["pairs"] + pairs)[-200:]
    save_json(f, d)
